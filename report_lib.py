"""Rendering library for the BuildConnect project report (Word .docx).

Faithfully mirrors the sample report template:
- A4, Times New Roman 12pt body, 1.5 line spacing, justified, black & white
- Chapter headings 16pt bold centered, section headings 14pt bold, sub-headings 12pt bold
- Bullets: manual bullet char, 0.25" -> 0.5" hanging indent
- Footer page numbers starting from Chapter 1 (front matter unnumbered)
- Figure placeholders (bordered empty box + caption)
"""

from docx import Document
from docx.shared import Pt, Inches, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
)
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
# A4 width 8.2677" - 1" left margin - 0.45" right margin
TEXT_WIDTH = 6.8177


# ---------------------------------------------------------------- low level

def _set_rfonts(rpr_owner, name=FONT):
    """Force explicit font attributes (kills theme-font inheritance)."""
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)):
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    _set_rfonts(normal.element)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def setup_section(sec, top=0.9, bottom=1.0, left=1.0, right=0.45, footer=0.65):
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Inches(top)
    sec.bottom_margin = Inches(bottom)
    sec.left_margin = Inches(left)
    sec.right_margin = Inches(right)
    sec.footer_distance = Inches(footer)


def add_page_field(paragraph):
    """Footer PAGE number field (12pt TNR)."""
    run = paragraph.add_run()
    run.font.name = FONT
    run.font.size = Pt(12)
    _set_rfonts(run._r)
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " PAGE "
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)


def set_pgnum_start(sec, start=1):
    sectPr = sec._sectPr
    pgMar = sectPr.find(qn("w:pgMar"))
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:start"), str(start))
    pgMar.addnext(el)


def table_borders(table, sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)


def table_cell_margins(table, left=360, right=360, top=58, bottom=58):
    """twips (1440 = 1")."""
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def table_fixed_layout(table):
    tblPr = table._tbl.tblPr
    el = OxmlElement("w:tblLayout")
    el.set(qn("w:type"), "fixed")
    tblPr.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def para_border_bottom(p, sz=6):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "000000")
    pBdr.append(b)
    pPr.append(pBdr)


# ------------------------------------------------------------- paragraph API

def _fmt(p, align=None, space_before=None, space_after=None, line=1.5,
         left=None, first=None, right=None):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if line is not None:
        pf.line_spacing = line
    if left is not None:
        pf.left_indent = Inches(left)
    if first is not None:
        pf.first_line_indent = Inches(first)
    if right is not None:
        pf.right_indent = Inches(right)
    return p


def add_text(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False,
             size=12, space_before=0, space_after=0, line=1.5, left=None,
             first=None, right=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = FONT
    _set_rfonts(r._r)
    return _fmt(p, align, space_before, space_after, line, left, first, right)


def add_segs(doc, segs, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12,
             space_before=0, space_after=0, left=None, right=None):
    """segs: list of (text, style) where style in '', 'b', 'i', 'bi'."""
    p = doc.add_paragraph()
    for text, style in segs:
        r = p.add_run(text)
        r.bold = "b" in style
        r.italic = "i" in style
        r.font.size = Pt(size)
        r.font.name = FONT
        _set_rfonts(r._r)
    return _fmt(p, align, space_before, space_after, 1.5, left, None, right)


def add_bullet(doc, text=None, lead=None, rest=None, space_after=4):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.25)
    pf.tab_stops.add_tab_stop(Inches(0.5))
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r0 = p.add_run("\u2022\t")
    r0.font.name = FONT
    r0.font.size = Pt(12)
    _set_rfonts(r0._r)
    if lead is not None:
        rb = p.add_run(lead)
        rb.bold = True
        rb.font.name = FONT
        rb.font.size = Pt(12)
        _set_rfonts(rb._r)
        body = rest if rest is not None else ""
    else:
        body = text if text is not None else ""
    if body:
        if lead is not None and not body.startswith(" "):
            body = " " + body
        r = p.add_run(body)
        r.font.name = FONT
        r.font.size = Pt(12)
        _set_rfonts(r._r)
    return p


def add_table(doc, headers, rows, widths, header_bold=True, pad_left=360,
              pad_right=None, borders=True):
    """3-col style data table. widths in inches (sum ~= TEXT_WIDTH)."""
    ncols = len(headers)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_fixed_layout(table)
    if borders:
        table_borders(table)
    table_cell_margins(table, left=pad_left,
                       right=pad_right if pad_right is not None else pad_left)
    repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(h)
        r.bold = header_bold
        r.font.name = FONT
        r.font.size = Pt(12)
        _set_rfonts(r._r)
    for row_vals in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row_vals):
            cell = cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            lines = str(val).split("\n")
            for j, line in enumerate(lines):
                p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                r = p.add_run(line)
                r.font.name = FONT
                r.font.size = Pt(12)
                _set_rfonts(r._r)
    for i, w in enumerate(widths):
        for row in table.rows:
            row.cells[i].width = Inches(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.line_spacing = 1.0
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    for r in spacer.runs:
        r.font.size = Pt(6)
    return table


def add_placeholder(doc, height_in, label="[ Insert image here ]", caption=None,
                    cap_size=12, cap_italic=True, cap_before=4, note=None):
    """Bordered empty box (image placeholder) + optional caption + note."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_fixed_layout(table)
    table_borders(table, sz=6)
    table_cell_margins(table, left=144, right=144, top=144, bottom=144)
    row = table.rows[0]
    row.height = Inches(height_in)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    cell.width = Inches(TEXT_WIDTH)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(label)
    r.italic = True
    r.font.name = FONT
    r.font.size = Pt(12)
    _set_rfonts(r._r)
    if caption:
        cap_p = add_text(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=cap_size,
                         italic=cap_italic, space_before=cap_before, space_after=4)
        if note:
            cap_p.paragraph_format.keep_with_next = True
    if note:
        add_text(doc, note, space_after=0)
    else:
        sp = doc.add_paragraph()
        sp.paragraph_format.line_spacing = 1.0
        sp.paragraph_format.space_after = Pt(0)
    # keep the box glued to its caption (but never chain whole figure groups)
    if caption:
        cell.paragraphs[-1].paragraph_format.keep_with_next = True
    return table


# ------------------------------------------------------------------ cover

def build_cover(doc, c):
    add_text(doc, c["title"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=20,
             space_before=15, space_after=10)
    add_text(doc, c["subtitle"], align=WD_ALIGN_PARAGRAPH.CENTER, italic=True,
             size=14, space_after=14)
    add_text(doc, c["course"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14,
             space_after=8)
    add_text(doc, "Submitted by", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             italic=True, size=13, space_after=6)
    add_text(doc, c["name_reg"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=15, space_after=14)
    add_text(doc, "In partial fulfillment for the award of the degree",
             align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12, space_after=0)
    add_text(doc, "of", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12,
             space_after=4)
    add_text(doc, "BACHELOR OF ENGINEERING", align=WD_ALIGN_PARAGRAPH.CENTER,
             bold=True, italic=True, size=14, space_after=6)
    add_text(doc, "in", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True,
             size=12, space_after=4)
    add_text(doc, "COMPUTER SCIENCE AND ENGINEERING",
             align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True, size=14,
             space_after=0)
    add_text(doc, c["college"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=14, space_before=248, space_after=6)
    add_text(doc, c["affiliation"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=12, space_after=4)
    add_text(doc, c["address"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=12, space_after=10)
    add_text(doc, c["date"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)


# -------------------------------------------------------------- bonafide

def build_bonafide(doc, c):
    add_text(doc, c["college"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=14, space_before=0, space_after=4)
    add_text(doc, c["affiliation"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=12, space_after=2)
    add_text(doc, c["address"], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
             size=12, space_after=0)
    add_text(doc, "BONAFIDE CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER,
             bold=True, size=15, space_before=150, space_after=24)
    add_text(doc, c["bonafide_line1"], size=12, space_after=0)
    add_text(doc, c["bonafide_line2"], size=12, space_after=0)
    # guide signature (right side)
    p = add_text(doc, "", size=12, space_before=150, space_after=0,
                 left=3.4, right=0.05)
    para_border_bottom(p)
    add_text(doc, "Course Instructor / Guide", align=WD_ALIGN_PARAGRAPH.RIGHT,
             size=12, space_before=2, space_after=40)
    # viva-voce lines
    add_segs(doc, [
        ("Submitted to the ", ""),
        (c["course"], "b"),
        (" Viva-Voce examination", ""),
    ], size=12, space_after=0)
    add_text(doc, c["viva_line"], size=12, space_after=0)
    # examiner signatures (two columns)
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    table_fixed_layout(t)
    table_cell_margins(t, left=0, right=0, top=20, bottom=20)
    widths = [TEXT_WIDTH / 2, TEXT_WIDTH / 2]
    for i, w in enumerate(widths):
        for row in t.rows:
            row.cells[i].width = Inches(w)
    # line row
    for col, indents in ((0, (0.0, 0.9)), (1, (0.9, 0.0))):
        cell = t.rows[0].cells[col]
        cell.width = Inches(widths[col])
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(40)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(indents[0])
        p.paragraph_format.right_indent = Inches(indents[1])
        p.add_run(" ")
        para_border_bottom(p)
    labels = t.rows[1].cells
    for col, (txt, align) in enumerate([
        ("Internal Examiner", WD_ALIGN_PARAGRAPH.LEFT),
        ("Co Examiner", WD_ALIGN_PARAGRAPH.RIGHT),
    ]):
        labels[col].width = Inches(widths[col])
        p = labels[col].paragraphs[0]
        p.alignment = align
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(txt)
        r.bold = True
        r.font.name = FONT
        r.font.size = Pt(12)
        _set_rfonts(r._r)


# ------------------------------------------------------------------- TOC

def build_toc(doc, title, rows):
    """rows: list of (ch_no, title, page_no). Borderless table like sample."""
    add_text(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12,
             space_after=0)
    widths = [0.9, 4.05, 1.87]
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    table_fixed_layout(t)
    table_cell_margins(t, left=100, right=100, top=40, bottom=40)

    def fill(cell, text, size, bold, align):
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold = bold
        r.font.name = FONT
        r.font.size = Pt(size)
        _set_rfonts(r._r)

    hdr = t.rows[0]
    for i, txt in enumerate(("CH NO", "TITLE", "PAGE NO")):
        hdr.cells[i].width = Inches(widths[i])
        fill(hdr.cells[i], txt, 14, True, WD_ALIGN_PARAGRAPH.CENTER)
    for ch, ttl, pg in rows:
        cells = t.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        cells[2].width = Inches(widths[2])
        fill(cells[0], ch, 12, False, WD_ALIGN_PARAGRAPH.CENTER)
        fill(cells[1], ttl, 12, False, WD_ALIGN_PARAGRAPH.LEFT)
        fill(cells[2], pg, 12, False, WD_ALIGN_PARAGRAPH.CENTER)


# --------------------------------------------------------------- renderer

def render(doc, blocks, chapter_break=False):
    for b in blocks:
        kind = b[0]
        if kind == "chapter":
            p1 = add_text(doc, b[1], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
                          size=16, space_before=0, space_after=0)
            p1.paragraph_format.keep_with_next = True
            if chapter_break:
                p1.paragraph_format.page_break_before = True
            p2 = add_text(doc, b[2], align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
                          size=16, space_before=0, space_after=4)
            p2.paragraph_format.keep_with_next = True
        elif kind == "h2":
            p = add_text(doc, b[1], align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=14,
                         space_before=10, space_after=6)
            p.paragraph_format.keep_with_next = True
        elif kind == "h3":
            p = add_text(doc, b[1], align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12,
                         space_before=8, space_after=4)
            p.paragraph_format.keep_with_next = True
        elif kind == "p":
            add_text(doc, b[1])
        elif kind == "psegs":
            add_segs(doc, b[1])
        elif kind == "bullet":
            add_bullet(doc, text=b[1])
        elif kind == "bulletr":
            add_bullet(doc, lead=b[1], rest=b[2])
        elif kind == "table":
            _, headers, rows, widths = b[:4]
            pad = b[4] if len(b) > 4 else 360
            add_table(doc, headers, rows, widths, pad_left=pad)
        elif kind == "figbox":
            _, height, caption = b[:3]
            cap_size = b[3] if len(b) > 3 else 12
            cap_ital = b[4] if len(b) > 4 else True
            note = b[5] if len(b) > 5 else None
            label = b[6] if len(b) > 6 else "[ Insert image here ]"
            add_placeholder(doc, height, label=label, caption=caption,
                            cap_size=cap_size, cap_italic=cap_ital, note=note)
        elif kind == "box":
            add_placeholder(doc, b[1], label=b[2])
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "gap":
            add_text(doc, "", space_after=b[1], size=12)
        elif kind == "fnote":
            add_text(doc, b[1], italic=True, size=12)
        else:
            raise ValueError(f"unknown block: {kind}")


def new_doc():
    doc = Document()
    setup_styles(doc)
    setup_section(doc.sections[0])
    return doc
